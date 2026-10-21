"""Extract external 2021 human-burden reference cases without scoring them.

The source is WHO's downloadable GHE 2021 global DALY workbook. Its cause
categories are not TULIP nodes and the rows here are not proposed score anchors.
"""

import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DIRECTORY = ROOT / "artifacts/tulip-consequence-rule-audit-2026-09-27"
SOURCE = DIRECTORY / "who-ghe2021-daly-global.xlsx"
OUTPUT = DIRECTORY / "who-ghe2021-daly-benchmark-cases.json"
SOURCE_SHA256 = "483ac1ae68d9191d9f739c2493cd8aaf2a92e87aecf0522ff7655a1d08501de7"
SOURCE_URL = (
    "https://cdn.who.int/media/docs/default-source/gho-documents/"
    "global-health-estimates/ghe2021_daly_global_new.xlsx?sfvrsn=cbefe871_3"
)
# WHO cause codes and labels are checked against the workbook, so row position
# alone cannot silently select a different condition after a source update.
SELECTED = {
    365: "Leprosy",
    310: "Trachoma",
    300: "Dengue",
    1570: "Drowning",
    220: "Malaria",
    30: "Tuberculosis",
    1530: "Road injury",
    390: "Lower respiratory infections",
    395: "COVID-19",
    0: "All Causes",
}


def main():
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    if digest != SOURCE_SHA256:
        raise ValueError(f"WHO source workbook changed: {digest}")
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook["Global 2021"]
    if sheet["F3"].value != "Global" or sheet["F4"].value != 2021:
        raise ValueError("WHO sheet geography or year changed")
    cases = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=10, values_only=True), start=10):
        code = row[0]
        if code not in SELECTED:
            continue
        label = next((item for item in reversed(row[2:6]) if isinstance(item, str) and item.strip()), None)
        if label != SELECTED[code]:
            raise ValueError(f"WHO cause {code} changed: {label!r}")
        value = row[6]
        if not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
            raise ValueError(f"WHO burden for {label} is invalid")
        cases.append({
            "causeCode": code,
            "cause": label,
            "modeledDalys": value,
            "unit": "DALY",
            "year": 2021,
            "geography": "global",
            "sheet": sheet.title,
            "cell": f"G{row_number}",
            "role": "external_burden_magnitude_context_only",
        })
    if {case["causeCode"] for case in cases} != set(SELECTED):
        raise ValueError("Selected WHO cause rows are incomplete")
    cases.sort(key=lambda case: case["modeledDalys"])
    report = {
        "version": "who_ghe2021_external_daly_benchmark_cases_2026_09_27_v1",
        "status": "verified_source_rows_no_score_anchors_selected",
        "sourcePage": "https://www.who.int/data/gho/data/themes/mortality-and-global-health-estimates/global-health-estimates-leading-causes-of-dalys/",
        "sourceWorkbook": SOURCE_URL,
        "sourceWorkbookSha256": digest,
        "sourceLocator": "WHO GHE 2021 Global summary DALY workbook, Global 2021 sheet, both sexes and all ages, column G; row cells listed per case.",
        "interpretation": "WHO cause-specific annual health loss provides an external magnitude ruler for considering 1–10 anchors. These are not climate-attributed impacts, some categories sit within broader categories, and All Causes contains the others. Do not add these rows or use their rank/percentiles to set TULIP anchors. They have no TULIP score or calibration status.",
        "cases": cases,
        "anchorAssignments": None,
        "publicReady": False,
    }
    OUTPUT.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps({"status": report["status"], "cases": len(cases),
        "smallest": cases[0]["modeledDalys"], "largest": cases[-1]["modeledDalys"]}, indent=2))


if __name__ == "__main__":
    main()
