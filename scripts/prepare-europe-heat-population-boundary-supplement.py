#!/usr/bin/env python3
"""Prepare explicit boundary/date supplements for the European heat population panel."""

from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / "tmp" / "tulip-regional-pattern"
ARTIFACT_DIR = ROOT / "artifacts" / "tulip-calibration-2026-09-22"
OUTPUT = ARTIFACT_DIR / "europe-heat-population-boundary-supplement.json"
CHECKED_ON = "2026-09-23"
SOURCES = {
    "nuts_changes": {
        "filename": "NUTS2021.xlsx",
        "url": "https://ec.europa.eu/eurostat/documents/345175/629341/NUTS2021.xlsx",
        "sha256": "b17dcc379bb3586550ec3b16f3f474dd2a8a55b4cfdcdc1e1575097f1c2d4761",
    },
    "ons_mid_2022": {
        "filename": "ons-mye22final.xlsx",
        "url": "https://www.ons.gov.uk/file?uri=%2Fpeoplepopulationandcommunity%2Fpopulationandmigration%2Fpopulationestimates%2Fdatasets%2Fpopulationestimatesforukenglandandwalesscotlandandnorthernireland%2Fmid2022%2Fmye22final.xlsx",
        "sha256": "89211116c238425a68b098b900d80b4f2f032578cf9c0d7cf5b38b20c96b6fcb",
    },
    "lau_2019_nuts_2016": {
        "filename": "EU-28-LAU-2019-NUTS-2016.xlsx",
        "url": "https://ec.europa.eu/eurostat/documents/345175/501971/EU-28-LAU-2019-NUTS-2016.xlsx",
        "sha256": "943510ed416e470e95919ebcf0d524f5cb0bb73fa5af8b4658f1ae8acb957df8",
    },
    "lau_2022_nuts_2021": {
        "filename": "EU-27-LAU-2022-NUTS-2021.xlsx",
        "url": "https://ec.europa.eu/eurostat/documents/345175/501971/EU-27-LAU-2022-NUTS-2021.xlsx",
        "sha256": "f1d26b6b2040d9dbceb255810f751e685081f9051fcbe82167d4a399da0df5dc",
    },
}
UK_CODE_MAP = {
    "UKC": "E12000001",
    "UKD": "E12000002",
    "UKE": "E12000003",
    "UKF": "E12000004",
    "UKG": "E12000005",
    "UKH": "E12000006",
    "UKI": "E12000007",
    "UKJ": "E12000008",
    "UKK": "E12000009",
    "UKL": "W92000004",
    "UKM": "S92000003",
    "UKN": "N92000002",
}
CHANGED_BE_EE = {
    "BE221", "BE222", "BE321", "BE322", "BE324", "BE325", "BE326", "BE327", "EE006", "EE007"
}
MANUAL_MERGED_LAU = {
    ("BE", "Opglabbeek"): "72042",
    ("BE", "Meeuwen-Gruitrode"): "72042",
    ("BE", "Neerpelt"): "72043",
    ("BE", "Overpelt"): "72043",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def pinned_source(source: dict[str, str]) -> Path:
    path = CACHE / source["filename"]
    if not path.exists():
        CACHE.mkdir(parents=True, exist_ok=True)
        request = Request(source["url"], headers={"User-Agent": "LostPlanet-TULIP/1.0 (+population-boundary-supplement)"})
        with urlopen(request, timeout=180) as response, path.open("wb") as output:
            output.write(response.read())
    actual = sha256(path)
    if actual != source["sha256"]:
        raise RuntimeError(f"{path.name}: source checksum changed: {actual}")
    return path


nuts_path = pinned_source(SOURCES["nuts_changes"])
ons_path = pinned_source(SOURCES["ons_mid_2022"])
lau_2019_path = pinned_source(SOURCES["lau_2019_nuts_2016"])
lau_2022_path = pinned_source(SOURCES["lau_2022_nuts_2021"])
population_panel = json.loads((ARTIFACT_DIR / "europe-heat-population-panel.json").read_text())

nuts_workbook = load_workbook(nuts_path, read_only=True, data_only=True)
nuts_sheet = nuts_workbook["Changes NUTS-3"]
nuts_changes = []
for row in nuts_sheet.iter_rows(min_row=2, values_only=True):
    old_code, new_code, label, change, explanation = row[:5]
    if old_code in CHANGED_BE_EE or new_code in CHANGED_BE_EE:
        nuts_changes.append({
            "oldCode": old_code,
            "newCode": new_code,
            "label": label,
            "change": change,
            "explanation": explanation,
        })
old_codes_documented = {row["oldCode"] for row in nuts_changes if row["oldCode"] in CHANGED_BE_EE}
if old_codes_documented != CHANGED_BE_EE:
    raise RuntimeError(f"NUTS workbook did not document every changed Belgian/Estonian code: {sorted(CHANGED_BE_EE - old_codes_documented)}")

ons_workbook = load_workbook(ons_path, read_only=True, data_only=True)
ons_sheet = ons_workbook["MYE2 - Persons"]
ons_rows = {}
for row in ons_sheet.iter_rows(min_row=9, values_only=True):
    code, name, geography, all_ages = row[:4]
    if code in UK_CODE_MAP.values():
        ons_rows[code] = {"name": name, "geography": geography, "population": int(all_ages)}
if set(ons_rows) != set(UK_CODE_MAP.values()):
    raise RuntimeError("ONS workbook did not contain all twelve required UK region/country rows")

uk_values = [{
    "studyRegionId": study_id,
    "onsCode": ons_code,
    **ons_rows[ons_code],
} for study_id, ons_code in UK_CODE_MAP.items()]
uk_total = sum(row["population"] for row in uk_values)
if uk_total != 67_596_281:
    raise RuntimeError(f"ONS regional values do not sum to the published UK total: {uk_total}")


def normalise(value: object) -> str:
    return str(value or "").strip().casefold()


lau_2019_workbook = load_workbook(lau_2019_path, read_only=True, data_only=True)
lau_2022_workbook = load_workbook(lau_2022_path, read_only=True, data_only=True)
matched_lau_groups: dict[tuple[str, str], list[tuple[tuple, tuple, str]]] = defaultdict(list)
old_lau_rows = 0
match_counts = defaultdict(int)
for country in ["BE", "EE"]:
    old_rows = [
        row for row in lau_2019_workbook[country].iter_rows(min_row=2, values_only=True)
        if row[0] in CHANGED_BE_EE
    ]
    new_rows = [
        row for row in lau_2022_workbook[country].iter_rows(min_row=2, values_only=True)
        if row[0] and str(row[0]).startswith(country)
    ]
    new_by_code = {str(row[1]).strip(): row for row in new_rows}
    old_lau_rows += len(old_rows)
    for old_row in old_rows:
        old_code = str(old_row[1]).strip()
        match = new_by_code.get(old_code)
        match_method = "same_lau_code"
        if match is None:
            old_names = {normalise(old_row[2]), normalise(old_row[3])} - {""}
            name_matches = [
                row for row in new_rows
                if old_names & ({normalise(row[3]), normalise(row[4])} - {""})
            ]
            if len(name_matches) == 1:
                match = name_matches[0]
                match_method = "same_name_changed_lau_code"
        if match is None:
            merged_code = MANUAL_MERGED_LAU.get((country, str(old_row[2])))
            if merged_code:
                match = new_by_code[merged_code]
                match_method = "documented_municipality_merger"
        if match is None:
            raise RuntimeError(f"No 2022 LAU match for {country} {old_row[:4]}")
        match_counts[match_method] += 1
        matched_lau_groups[(country, str(match[1]).strip())].append((old_row, match, match_method))

if old_lau_rows != 110 or len(matched_lau_groups) != 108:
    raise RuntimeError(f"Unexpected LAU reconstruction coverage: {old_lau_rows} old rows, {len(matched_lau_groups)} new rows")

region_totals = defaultdict(float)
merge_allocations = []
for (country, new_lau_code), matches in matched_lau_groups.items():
    new_row = matches[0][1]
    new_population = float(new_row[6])
    old_region_ids = {match[0][0] for match in matches}
    if len(old_region_ids) == 1:
        region_totals[next(iter(old_region_ids))] += new_population
        continue
    old_population_total = sum(float(match[0][5]) for match in matches)
    allocations = defaultdict(float)
    for old_row, _, _ in matches:
        allocations[old_row[0]] += new_population * float(old_row[5]) / old_population_total
    for region_id, allocated_population in allocations.items():
        region_totals[region_id] += allocated_population
    merge_allocations.append({
        "country": country,
        "newLauCode": new_lau_code,
        "newLauName": new_row[3],
        "population2022": int(new_population),
        "oldComponents": [{
            "oldRegionId": old_row[0],
            "oldLauCode": str(old_row[1]),
            "oldLauName": old_row[2],
            "population2019": int(old_row[5]),
            "allocationShare": float(old_row[5]) / old_population_total,
        } for old_row, _, _ in matches],
        "rule": "Allocate the merged municipality's 2022 total using its two predecessor municipalities' 2019 population shares.",
    })

if set(region_totals) != CHANGED_BE_EE or len(merge_allocations) != 1:
    raise RuntimeError("Old-boundary population reconstruction did not produce the expected ten regions and one cross-boundary merge")
reconstructed_regions = [{
    "regionId": region_id,
    "population2022": round(region_totals[region_id], 3),
    "status": "proportional_merge_sensitivity" if region_id in {"BE221", "BE222"} else "lau_aggregate_reconstructed",
} for region_id in sorted(region_totals)]
reconstructed_total = round(sum(region_totals.values()), 3)

panel_by_id = {row["id"]: row for row in population_panel["regions"]}
nearest_total = population_panel["checks"]["latestNearestYearCandidateTotal"]
nearest_uk_total = sum(panel_by_id[study_id]["nearestObservedCandidateByYear"]["2022"]["value"] for study_id in UK_CODE_MAP)
nearest_be_ee_total = sum(panel_by_id[region_id]["nearestObservedCandidateByYear"]["2022"]["value"] for region_id in CHANGED_BE_EE)
uk_supplemented_sensitivity_total = nearest_total - nearest_uk_total + uk_total
fully_supplemented_sensitivity_total = round(uk_supplemented_sensitivity_total - nearest_be_ee_total + reconstructed_total, 3)
published_coverage_minimum = population_panel["checks"]["publishedCoverageFloor"]

payload = {
    "version": "tulip_europe_heat_population_boundary_supplement_2026_09_23_v3",
    "generatedOn": CHECKED_ON,
    "status": "official_uk_midyear_and_lau_2022_edition_boundary_values_ingested_lau_reference_date_and_two_region_split_remain_sensitivity",
    "productionRatingsEnabled": 0,
    "purpose": "Separate geography-compatible sensitivity inputs from exact boundary-aligned January population observations instead of silently carrying old values forward or overstating Eurostat validation.",
    "sources": {
        source_id: {
            "url": source["url"],
            "localPath": str((CACHE / source["filename"]).relative_to(ROOT)),
            "bytes": (CACHE / source["filename"]).stat().st_size,
            "sha256": source["sha256"],
        }
        for source_id, source in SOURCES.items()
    },
    "ukMid2022": {
        "status": "official_geography_compatible_midyear_sensitivity_values",
        "referenceDate": "2022-06-30",
        "targetEurostatReferenceDate": "2022-01-01",
        "dateAlignment": "not_exact",
        "regions": uk_values,
        "regionCount": len(uk_values),
        "total": uk_total,
        "rule": "These values may test population-weighting sensitivity. They cannot be labelled as the article's Eurostat January observations without an explicit date-alignment method.",
    },
    "belgiumEstoniaBoundaryChanges": {
        "status": "eight_old_boundary_regions_reconstructed_two_require_proportional_merge_split",
        "oldRegionCount": len(CHANGED_BE_EE),
        "oldRegionIds": sorted(CHANGED_BE_EE),
        "changeRecords": nuts_changes,
        "oldLauRows": old_lau_rows,
        "matchedNewLauRows": len(matched_lau_groups),
        "matchCounts": dict(sorted(match_counts.items())),
        "regions": reconstructed_regions,
        "reconstructedTotal": reconstructed_total,
        "directLauAggregateRegions": 8,
        "proportionalMergeSensitivityRegions": 2,
        "mergeAllocations": merge_allocations,
        "lauPopulationField": {
            "edition": "LAU 2022",
            "field": "POPULATION",
            "providedInEurostatTransmissionWorkbook": True,
            "voluntary": True,
            "validatedByEurostat": False,
            "referenceDate": "not specified in workbook",
            "workbookEvidence": "The Overview sheet labels population and area as voluntary and not validated by Eurostat.",
            "usage": "Boundary-reconstruction sensitivity only; not an exact Eurostat 1 January 2022 observation.",
        },
        "rule": "Aggregate the LAU 2022 edition population field using each LAU's 2019 NUTS 2016 membership. Eurostat does not validate this voluntary field and the workbook does not state its exact reference date. Only the Oudsbergen merger crosses two old regions; its split remains a declared sensitivity based on 2019 predecessor population shares.",
    },
    "checks": {
        "originalMissing2022Regions": population_panel["checks"]["latestYearMissingRegions"],
        "officialUkSensitivityRegions": len(uk_values),
        "boundaryChangedRegionsReconstructed": len(CHANGED_BE_EE),
        "boundaryChangedDirectLauAggregateRegions": 8,
        "boundaryChangedProportionalMergeSensitivityRegions": 2,
        "originalNearestYearSensitivityTotal": nearest_total,
        "originalNearestYearUkTotal": nearest_uk_total,
        "officialUkMidyearTotal": uk_total,
        "originalNearestYearBelgiumEstoniaTotal": nearest_be_ee_total,
        "reconstructedBelgiumEstoniaTotal": reconstructed_total,
        "ukSupplementedSensitivityTotal": uk_supplemented_sensitivity_total,
        "fullySupplementedSensitivityTotal": fully_supplemented_sensitivity_total,
        "articleStatedCoverageMinimum": published_coverage_minimum,
        "differenceFromArticleStatedCoverageMinimum": fully_supplemented_sensitivity_total - published_coverage_minimum,
    },
    "interpretation": {
        "gain": "Official 2022 population values now exist for all twelve study-era UK NUTS1 slots, with an explicit six-month reference-date difference.",
        "limit": "Eight Belgian and Estonian values are direct old-boundary aggregates of the LAU 2022 edition population field, which is voluntary, not validated by Eurostat and has no exact reference date stated in the workbook. BE221 and BE222 also share the later Oudsbergen municipality, so their individual values depend on one declared proportional split.",
        "next": "Obtain date-defined national-statistics population values for the ten Belgian and Estonian old-boundary regions, replace the Oudsbergen proportional split, decide the UK date-alignment rule, then compare the resulting total and weighted warming mean with the article.",
    },
}
payload["digest"] = hashlib.sha256(json.dumps({
    "sources": payload["sources"],
    "ukMid2022": payload["ukMid2022"],
    "belgiumEstoniaBoundaryChanges": payload["belgiumEstoniaBoundaryChanges"],
    "checks": payload["checks"],
}, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
OUTPUT.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n")
print(json.dumps({"output": str(OUTPUT.relative_to(ROOT)), "status": payload["status"], "checks": payload["checks"], "digest": payload["digest"]}, indent=2))
