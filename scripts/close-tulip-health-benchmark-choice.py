"""Reproduce fixed health references from the WHO workbook, not node scores.

The additional WHO rows were not used to select the four middle thresholds.
They challenge magnitude coverage; they are not an independent causal model,
climate-attributed node assessments, or empirical proof of the policy weights.
"""
import hashlib
import json
from decimal import Decimal, ROUND_HALF_UP, localcontext
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DIRECTORY = ROOT / "artifacts/tulip-calibration-closeout-2026-09-28"
HISTORICAL = ROOT / "artifacts/tulip-consequence-rule-audit-2026-09-27"


def map_quantity(value, points):
    with localcontext() as ctx:
        ctx.prec = 60
        x = Decimal(str(value))
        if x >= Decimal(str(points[-1]["burden"])):
            return Decimal(10)
        for lo, hi in zip(points, points[1:]):
            a, b = Decimal(str(lo["burden"])), Decimal(str(hi["burden"]))
            if x <= b:
                t = ((1 + x).ln() - (1 + a).ln()) / ((1 + b).ln() - (1 + a).ln())
                return Decimal(lo["score"]) + Decimal(hi["score"] - lo["score"]) * t
        raise ValueError("No fixed benchmark interval")


def main():
    anchor_path = ROOT / "docs/tulip-calibration/human-health-anchor-candidate.json"
    decision_path = ROOT / "docs/tulip-calibration/calibration-decisions-2026-09-28.json"
    anchors = json.loads(anchor_path.read_text())
    decision = json.loads(decision_path.read_text())
    source_path = HISTORICAL / "who-ghe2021-daly-global.xlsx"
    source_sha = hashlib.sha256(source_path.read_bytes()).hexdigest()
    assert source_sha == anchors["sourceWorkbookSha256"]
    assert decision["humanHealth"]["anchorVersion"] == anchors["version"]
    assert decision["publication"]["wholeMethodCalibrated"] is False
    initial = json.loads((HISTORICAL / "who-ghe2021-daly-benchmark-cases.json").read_text())
    original_codes = {row["causeCode"] for row in initial["cases"]}
    workbook = load_workbook(source_path, data_only=True, read_only=True)
    sheet = workbook["Global 2021"]
    assert sheet["F3"].value == "Global" and sheet["F4"].value == 2021
    seen = set()
    cases = []
    for number, row in enumerate(sheet.iter_rows(min_row=10, values_only=True), 10):
        code = row[0]
        if not isinstance(code, int):
            continue
        assert code not in seen
        seen.add(code)
        label = next(x for x in reversed(row[2:6]) if isinstance(x, str) and x.strip())
        value = row[6]
        assert isinstance(value, (int, float)) and Decimal(str(value)).is_finite() and value >= 0
        score = map_quantity(value, anchors["points"])
        cases.append({
            "causeCode": code, "cause": label, "sourceCell": f"Global 2021!G{number}",
            "modeledDalys": value, "unroundedComponentDiagnostic": str(score),
            "componentDiagnostic": float(score.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)),
            "role": "initial_reference" if code in original_codes else "additional_external_magnitude_challenge",
            "nodeRating": None,
        })
    all_causes = next(x for x in cases if x["causeCode"] == 0)
    assert all_causes["modeledDalys"] == anchors["points"][-1]["burden"]
    assert all_causes["sourceCell"] == anchors["points"][-1]["sourceCell"]
    ordered = sorted(cases, key=lambda x: x["modeledDalys"])
    assert all(Decimal(a["unroundedComponentDiagnostic"]) <= Decimal(b["unroundedComponentDiagnostic"])
               for a, b in zip(ordered, ordered[1:]))
    for point in anchors["points"]:
        assert map_quantity(point["burden"], anchors["points"]) == point["score"]
    # These source categories are NOT summed into a burden or asserted to be a
    # verified mutually exclusive partition. Preserve any source discrepancy.
    root_codes = [10, 600, 1510]
    root_rows = [next(x for x in cases if x["causeCode"] == c) for c in root_codes]
    root_sum = sum(Decimal(str(x["modeledDalys"])) for x in root_rows)
    observed_gap = Decimal(str(all_causes["modeledDalys"])) - root_sum
    output = {
        "version": "tulip_health_benchmark_choice_source_check_2026_09_28_v1",
        "status": "internal_health_benchmark_choice_checked_not_whole_method_calibration",
        "decisionVersion": decision["version"],
        "inputs": {
            str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in [anchor_path, decision_path, source_path]
        },
        "sourceUrl": anchors["sourceWorkbook"],
        "scope": "2021 global both-sex all-age WHO cause-specific annual DALY point estimates",
        "summary": {"referenceRows": len(cases), "initialReferences": len(original_codes),
                    "additionalMagnitudeChallenges": sum(x["role"] != "initial_reference" for x in cases),
                    "wholeNodeAssessments": 0},
        "selectedBenchmarks": anchors["points"], "cases": ordered,
        "rootCategoryCheck": {
            "causeCodes": root_codes, "sourceCells": [x["sourceCell"] for x in root_rows],
            "sum": str(root_sum), "allCauses": str(all_causes["modeledDalys"]),
            "allCausesMinusSelectedRootSum": str(observed_gap),
            "interpretation": "An observed source-table difference, not a repaired total. Cause rows are contextual cases; no category sum is used for the anchors or an attributable node. The reason for the difference has not been established.",
        },
        "checks": {"sourceWorkbookDigestMatched": True, "outerReferenceExactCellMatched": True,
                   "allMagnitudeOrderingsPreserved": True, "fixedBenchmarksReproduced": True,
                   "catalogueScoresUsedToFitBenchmarks": False},
        "limitations": [
            "The additional rows challenge scale coverage after the initial anchor proposal. They are not a blind or independent scientific review.",
            "These cause categories overlap hierarchically, are not climate-attributed observations, and are never added together as TULIP burden.",
            "Point estimates without source intervals are not zero-width uncertainty intervals.",
            "The four middle amounts and interpolation remain explicit TULIP policy choices.",
            "This check does not establish non-health human conversions, nature calibration or any complete node rating."
        ],
        "publicReady": False,
    }
    DIRECTORY.mkdir(parents=True, exist_ok=True)
    (DIRECTORY / "health-benchmark-source-check.json").write_text(json.dumps(output, indent=2) + "\n")
    print(json.dumps({"status": output["status"], **output["summary"],
                      "rootCategoryDifferencePreserved": str(observed_gap)}))


if __name__ == "__main__":
    main()
