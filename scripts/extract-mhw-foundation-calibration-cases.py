"""Extract typed historical reference cases from a published MHW synthesis.

These records are calibration evidence, never a global current harm estimate.
Run with a Python runtime containing openpyxl and supply the two source XLSX
files downloaded from the publisher. Their hashes and table headers are checked.
"""
import argparse
import hashlib
import json
import math
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts/tulip-consequence-rule-audit-2026-09-27'
PREFIX = 'https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-024-49307-9/MediaObjects/'
SOURCES = {
    'regions': {'sha256': '818aa121c7047989218c0ffe1bf6d8cf69ec346bac597cd34614ad6202e7c311',
                'url': PREFIX + '41467_2024_49307_MOESM4_ESM.xlsx',
                'cachedName': 'mhw-foundation-2024-supplementary-data-1.xlsx'},
    'figures': {'sha256': '4d2cd12e07177d2a808de186d2a5ce23179a71fc7dd89ca8727cfdeec55ac881',
                'url': PREFIX + '41467_2024_49307_MOESM7_ESM.xlsx',
                'cachedName': 'mhw-foundation-2024-source-data.xlsx'},
}
GROUPS = {
    'Hard Corals': ('hard_coral_bleaching_prevalence', 'stress_response_prevalence', 'inhibition_only', 1047, [1980, 2020], 'within_16_weeks_after_mhw'),
    'Gorgonian soft corals': ('gorgonian_mme_affected_population_prevalence', 'mortality_event_response_prevalence', 'inhibition_only', 76, [1985, 2020], 'within_16_weeks_after_mhw'),
    'Seagrass (change in cover)': ('seagrass_cover_relative_change', 'biological_abundance_state_change', 'signed_loss_and_gain', 60, [2005, 2022], '6_to_12_months_after_mhw'),
    'Macroalgae (change in cover)': ('macroalgal_cover_relative_change', 'biological_abundance_state_change', 'signed_loss_and_gain', 82, [1982, 2021], '6_to_12_months_after_mhw'),
    'Macroalgae (change in density)': ('macroalgal_density_relative_change', 'biological_abundance_state_change', 'signed_loss_and_gain', 57, [1982, 2021], '6_to_12_months_after_mhw'),
}


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def optional_number(value):
    if numeric(value):
        return value
    if value in (None, 'NA'):
        return None
    raise ValueError(f'Unexpected numeric field: {value!r}')


def source_workbook(path, kind):
    source = SOURCES[kind]
    data = path.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    if digest != source['sha256']:
        raise ValueError(f'{kind} source changed: {digest}')
    destination = OUT / source['cachedName']
    if destination.exists() and destination.read_bytes() != data:
        raise ValueError(f'Refusing to replace different cached source: {destination}')
    destination.write_bytes(data)
    return load_workbook(path, read_only=True, data_only=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('regions', type=Path)
    parser.add_argument('figures', type=Path)
    args = parser.parse_args()
    regional = source_workbook(args.regions, 'regions')['Sheet1']
    figures = source_workbook(args.figures, 'figures')
    headers = next(regional.iter_rows(values_only=True))
    if headers[:6] != ('Ecoregion', 'Foundation species group', 'n (total)', 'n (available as ranges)',
                       'n (available as accurate records)', 'Average change (% loss in red, % gain in green)'):
        raise ValueError('Regional source table headers changed')
    cases = []
    source_rows = list(regional.iter_rows(values_only=True))
    continuation_rows = set()
    for row_number, row in enumerate(source_rows[1:], start=2):
        if row_number in continuation_rows:
            continue
        if row[0] is None:
            if any(value is not None for value in row):
                raise ValueError('Partial blank source row')
            continue
        region, group, n, coarse, accurate, mean, response_range, sd = row
        group_fragments = [group]
        end_row = row_number
        # Two type labels wrap onto their own following spreadsheet row. They
        # specify density; treating the first line alone as an unknown metric
        # would misclassify twelve observations.
        if group == 'Macroalgae' and row_number < len(source_rows):
            continuation = source_rows[row_number]
            if continuation != (None, '(change in density)', None, None, None, None, None, None):
                raise ValueError(f'Unrecognized source group continuation at {row_number + 1}')
            group_fragments.append(continuation[1])
            group = 'Macroalgae (change in density)'
            continuation_rows.add(row_number + 1)
            end_row = row_number + 1
        if group not in GROUPS or not isinstance(n, int) or n <= 0 or not numeric(mean):
            raise ValueError(f'Invalid or unrecognized regional row {row_number}: {row!r}')
        metric, role, encoding, _, source_period, response_window = GROUPS[group]
        coarse, accurate, sd = optional_number(coarse), optional_number(accurate), optional_number(sd)
        if coarse is not None and accurate is not None and coarse + accurate != n:
            raise ValueError(f'Source observation counts do not reconcile at row {row_number}')
        cases.append({
            'caseId': f'mhw_foundation_2024_regional_row_{row_number}',
            'sourceLocator': f'Supplementary Data 1, Sheet1!A{row_number}:H{end_row}',
            'ecoregion': region.strip(), 'sourceGroup': group, 'endpointKind': metric,
            'sourceGroupFragments': group_fragments,
            'indicatorRole': role, 'directionEncoding': encoding,
            'sourceTableReportedMeanPercent': mean,
            'meanInhibitionPrevalenceMagnitudePercent': abs(mean) if encoding == 'inhibition_only' else None,
            'signedRelativeAbundanceChangePercent': mean if encoding == 'signed_loss_and_gain' else None,
            'observationCount': n, 'categoryCodedObservationCount': coarse,
            'accuratelyReportedObservationCount': accurate,
            'reportedAcrossObservationRange': response_range if response_range != 'NA' else None,
            'reportedAcrossObservationSd': sd,
            'sourceDatasetPeriodEnvelope': {'start': source_period[0], 'end': source_period[1],
                'basis': 'paper_table_2_group_dataset_envelope_not_row_assessment_period'},
            'postEventResponseWindow': response_window,
            'categoryCodingFlag': coarse is not None and coarse > 0,
            'positiveCoralSourceEncodingFlag': encoding == 'inhibition_only' and mean > 0,
            'singleObservationFlag': n == 1,
            'lossGainMixtureFlag': encoding == 'signed_loss_and_gain' and isinstance(response_range, str)
                and 'loss' in response_range and 'gain' in response_range,
            'coverage': 'historical_ecoregion_multi_event_reference',
            'grossPositiveLossMean': None,
            'attributableDamagedAreaKm2': None,
            'nodeAttributableFraction': None,
            'currentAssessmentPeriod': None,
            'score': None, 'publicReady': False,
        })
    grouped = Counter()
    for case in cases:
        grouped[case['sourceGroup']] += case['observationCount']
    if dict(grouped) != {group: values[3] for group, values in GROUPS.items()}:
        raise ValueError(f'Regional counts no longer reproduce the article: {grouped}')
    if len(cases) != 94 or sum(grouped.values()) != 1322:
        raise ValueError('Expected 94 regional summaries representing 1322 localized-area/event observations')
    # Published figure counts independently reproduce the group totals. Blank
    # cells here are count bins, not missing biological harm measurements.
    figure2 = list(figures['Figure 2'].iter_rows(values_only=True))
    figure3 = list(figures['Figure 3'].iter_rows(values_only=True))
    figure_counts = {
        'hard_corals': sum(value for row in figure2[2:] if row[0] == 'Hard Corals' for value in row[2:6] if numeric(value)),
        'gorgonians': sum(value for row in figure2[2:] if row[8] == 'Gorgonian soft corals' for value in row[10:13] if numeric(value)),
        'seagrass': sum(value for row in figure3[2:] if row[0] == 'Seagrass' for value in row[2:7] if numeric(value)),
        'macroalgae': sum(value for row in figure3[2:] if row[0] == 'macroalgae' for value in row[2:7] if numeric(value)),
    }
    if figure_counts != {'hard_corals': 1047, 'gorgonians': 76, 'seagrass': 60, 'macroalgae': 139}:
        raise ValueError(f'Figure counts disagree with regional source: {figure_counts}')
    figure_region_aliases = {
        'Oregon, Washington, Vancouver Coast and Shelf': 'Oregon, Washington, Vancouver Is.',
        'Tweed - Moreton': 'Tweed–Moreton',
    }
    plant_bins = {}
    for row in figure3[2:]:
        if row[0] not in ('Seagrass', 'macroalgae'):
            continue
        key = (row[0], row[1].strip())
        if key in plant_bins:
            raise ValueError(f'Duplicate figure group/region: {key}')
        counts = [0 if value is None else value for value in row[2:7]]
        if any(not isinstance(value, int) or value < 0 for value in counts):
            raise ValueError(f'Invalid figure count bins: {key}')
        plant_bins[key] = dict(zip(['severeLoss', 'moderateLoss', 'mildLossOrGain', 'moderateGain', 'severeGain'], counts))
    for case in cases:
        case['sourceResponseCategoryCounts'] = None
        case['observationsWithAtLeast10PercentLoss'] = None
        if case['directionEncoding'] != 'signed_loss_and_gain':
            continue
        group_key = 'Seagrass' if case['sourceGroup'].startswith('Seagrass') else 'macroalgae'
        region_key = figure_region_aliases.get(case['ecoregion'], case['ecoregion'])
        bins = plant_bins[(group_key, region_key)]
        if sum(bins.values()) != case['observationCount']:
            raise ValueError(f'Figure count does not match summary: {case["caseId"]}')
        case['sourceResponseCategoryCounts'] = bins
        case['observationsWithAtLeast10PercentLoss'] = bins['severeLoss'] + bins['moderateLoss']
    report = {
        'version': 'mhw_foundation_typed_reference_cases_2026_09_27_v1',
        'status': 'source_checked_historical_reference_pool_not_harm_calibration',
        'sourcePaper': 'https://doi.org/10.1038/s41467-024-49307-9',
        'sourcePaperVersion': 'Smith et al., Nature Communications 15:5052, published 13 June 2024',
        'sources': SOURCES,
        'sourcePeriodMetadataLocator': 'Article Table 2; group envelopes are editorial transcriptions of its dataset periods, not event dates for individual summary rows.',
        'cases': cases,
        'summary': {
            'regionalSummaryRows': len(cases), 'localizedAreaEventObservations': sum(grouped.values()),
            'observationsBySourceMetricGroup': dict(grouped), 'independentFigureCounts': figure_counts,
            'categoryCodedObservationCount': sum(case['categoryCodedObservationCount'] or 0 for case in cases),
            'singleObservationSummaryRows': sum(case['singleObservationFlag'] for case in cases),
            'positiveCoralSourceEncodingRows': sum(case['positiveCoralSourceEncodingFlag'] for case in cases),
            'lossGainMixtureRows': sum(case['lossGainMixtureFlag'] for case in cases),
            'plantSummaryRowsWithIndependentlyMatchedCountBins': len(plant_bins),
        },
        'limitations': [
            'The metric percentages are not interchangeable: bleaching prevalence is not coral mortality or ecosystem condition loss.',
            'Source abundance-change means can offset losses against gains. A mean, its SD and its observed range cannot recover gross positive loss or an uncertainty interval for attributable harm.',
            'Matched plant-response bins show sampled site/event counts, not area. The mild bin includes losses and gains; bin midpoints cannot recover continuous gross loss.',
            'Coral category-coded responses use conservative 1%, 11%, 51% replacements in the paper. They are not precise continuous observations; three single-category source means are positive-coded inhibition, not ecological gains.',
            'Counts refer to sampled localized area/event observations, not damaged habitat extent or a representative area-weighted global population.',
            'Rows combine historical events, response windows and sometimes different species. Publication year is not the biological assessment year.',
            'Temporal association with author-filtered strong summer marine heatwaves is useful evidence but is not a quantified no-heatwave causal counterfactual or a global climate-attributable share.',
            'The summary workbook cannot establish current persistence, a collapse threshold, whole-node coverage or a cross-ecosystem exchange rate.',
        ],
        'currentNodeAssessmentsCompleted': 0, 'score': None, 'publicReady': False,
    }
    (OUT / 'mhw-foundation-typed-reference-cases.json').write_text(json.dumps(report, indent=2, ensure_ascii=False) + '\n')
    print(json.dumps({'status': report['status'], 'summary': report['summary'], 'score': None}, indent=2))


if __name__ == '__main__':
    main()
