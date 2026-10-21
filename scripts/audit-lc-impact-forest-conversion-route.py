"""Test whether FAO's global deforestation rate can enter LC-IMPACT V2.

The diagnostic reads source-pinned country factors; it does not compute a
TULIP rating or multiply FAO's global rate by an incompatible country factor.
"""

import hashlib
import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DIRECTORY = ROOT / "artifacts/tulip-consequence-rule-audit-2026-09-27"
ARCHIVE = DIRECTORY / "lc-impact-v2-land-use.zip"
FAO_RECORD = DIRECTORY / "fao-fra-2025-global-deforestation-extent.json"
OUTPUT = DIRECTORY / "lc-impact-forest-conversion-route-audit.json"
ARCHIVE_SHA256 = "5c7dd62fc3530ba2e8a9289927229cb68df9ead3e3591d1ca9a8b18538dffa68"
ARCHIVE_URL = "https://zenodo.org/records/19202729/files/6_CF_LAND_USE.zip?download=1"
WORKBOOK = "CF_LAND_USE/EQ_LandUse_gloPDF_Country_Aggregated_2025-12-17.xlsx"
FLOW = "from forest, natural"
COUNTRIES = {"FIN": "Finland", "IDN": "Indonesia", "BRA": "Brazil", "USA": "United States of America"}


def main():
    actual_digest = hashlib.sha256(ARCHIVE.read_bytes()).hexdigest()
    if actual_digest != ARCHIVE_SHA256:
        raise ValueError(f"LC-IMPACT archive changed: {actual_digest}")
    fao = json.loads(FAO_RECORD.read_text())
    if fao["scope"]["geography"] != "global" or fao["scope"]["metricKind"] != "gross_deforestation":
        raise ValueError("FAO endpoint scope changed")
    with ZipFile(ARCHIVE) as bundle:
        sheet = load_workbook(BytesIO(bundle.read(WORKBOOK)), read_only=True, data_only=True)["Country"]
        columns = [cell.value for cell in sheet[1]]
        expected = ["FLOW_name", "LCIAMethod_location", "CF", "CF_unit", "FLOW_class1",
                    "Species", "LCIAMethod_realm", "LCIAMethod_mathematicalApproach"]
        if any(column not in columns for column in expected):
            raise ValueError("LC-IMPACT country factor columns changed")
        rows = []
        factor_diagnostics = {"countrySheetRows": 0, "negativeFactors": 0,
                              "missingFactors": 0, "units": {}}
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            record = dict(zip(columns, values))
            factor_diagnostics["countrySheetRows"] += 1
            unit = record["CF_unit"]
            factor_diagnostics["units"][unit] = factor_diagnostics["units"].get(unit, 0) + 1
            if record["CF"] is None:
                factor_diagnostics["missingFactors"] += 1
            elif isinstance(record["CF"], (int, float)) and record["CF"] < 0:
                factor_diagnostics["negativeFactors"] += 1
            if (record["FLOW_name"] == FLOW
                    and record["LCIAMethod_location"] in COUNTRIES
                    and record["FLOW_class1"] == "Land transformation"
                    and record["Species"] == "Aggregated"
                    and record["LCIAMethod_realm"] == "Terrestrial"
                    and record["LCIAMethod_mathematicalApproach"] == "Average"):
                if record["CF_unit"] != "Global PDF*y/m2" or not isinstance(record["CF"], (int, float)):
                    raise ValueError("LC-IMPACT factor unit or value changed")
                rows.append({
                    "countryCode": record["LCIAMethod_location"],
                    "country": COUNTRIES[record["LCIAMethod_location"]],
                    "flow": FLOW,
                    "factor": record["CF"],
                    "unit": record["CF_unit"],
                    "approach": "Average",
                    "sourceSheetCell": f"Country!G{row_number}",
                })
    if {row["countryCode"] for row in rows} != set(COUNTRIES) or len(rows) != len(COUNTRIES):
        raise ValueError("Expected one factor per selected country")
    if factor_diagnostics["negativeFactors"] <= 0 or factor_diagnostics["missingFactors"] <= 0:
        raise ValueError("Expected signed and missing source factors not reproduced")
    rows.sort(key=lambda row: row["factor"])
    by_country = {row["countryCode"]: row for row in rows}
    indonesia_to_finland_ratio = by_country["IDN"]["factor"] / by_country["FIN"]["factor"]
    if indonesia_to_finland_ratio <= 1000:
        raise ValueError("Expected LC-IMPACT location sensitivity not reproduced")
    result = {
        "version": "lc_impact_v2_forest_conversion_scope_audit_2026_09_27_v1",
        "status": "global_fao_rate_not_compatible_with_location_specific_factors",
        "landUseFactorSource": "https://zenodo.org/records/19202729",
        "archiveUrl": ARCHIVE_URL,
        "archiveSha256": actual_digest,
        "archiveWorkbook": WORKBOOK,
        "sourceSelection": "Country sheet; aggregated terrestrial; land transformation; average approach; from forest, natural; four illustrative country locations.",
        "factorRows": rows,
        "countrySheetFactorDiagnostics": factor_diagnostics,
        "indonesiaToFinlandFactorRatioForSameFlow": indonesia_to_finland_ratio,
        "faoSource": fao["sourceUrl"],
        "faoGlobalGrossDeforestationHectaresPerYear": fao["grossDeforestationRate"]["value"],
        "faoPeriod": {"start": fao["scope"]["periodStartYear"], "end": fao["scope"]["periodEndYear"]},
        "requiredToComputeACompatibleLandUseEndpoint": [
            "Location-specific converted forest area, not only a global mean annual total.",
            "Origin forest category and destination land-use category for each conversion, mapped to LC-IMPACT inventory flows.",
            "Consistent transformation/occupation accounting so recovery or occupied years are not multiplied twice.",
            "A matched annual assessment period, counterfactual, and exclusion of overlapping pathways.",
            "Sensitivity to country/ecoregion, average versus marginal characterization, model factors and missing forest types.",
            "Signed or missing characterization factors must retain their source meaning; do not turn them into zero harm or a positive danger score by default.",
        ],
        "interpretation": "The same named forest-origin flow differs greatly by modelled location. This is factor sensitivity, not an observed comparison of forest damage in those countries. FAO's 10.9 million hectares/year is global and does not supply the country, origin-to-destination and time accounting needed to calculate a PDF-year endpoint. Multiplying it by a selected factor would create a spurious number.",
        "ecosystemScope": "LC-IMPACT PDF-year represents modelled global species-loss risk and is not all ecosystem condition, service provision or local biodiversity loss. Terrestrial, freshwater and marine results need an explicit weighting choice. The country factor table also contains negative and missing values, so a positive-only score conversion cannot be applied blindly.",
        "natureScore": None,
        "publicReady": False,
    }
    OUTPUT.write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps({"status": result["status"], "factorRows": len(rows),
        "indonesiaToFinlandFactorRatio": round(indonesia_to_finland_ratio),
        "natureScore": None}, indent=2))


if __name__ == "__main__":
    main()
